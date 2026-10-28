# SPDX-FileCopyrightText: The Docling-pipeline Contributors
# SPDX-License-Identifier: MIT

"""Unit tests for password-protected document handling in the extract operator.

Covers the ``document_passwords`` configuration, encryption detection and in-memory
decryption, and how both extraction paths report documents that cannot be decrypted.
Encrypted fixtures are generated at test time so no secrets are checked in.
"""

import io
import logging
import sys
from typing import Any
from unittest.mock import MagicMock, patch

import pyarrow as pa
import pytest
from pydantic import BaseModel

from docpipe.core.constants.constants import DocpipeConstants, Metrics
from docpipe.core.constants.operator_constants import OperatorConstants
from docpipe.core.operators.extract import encrypted_documents
from docpipe.core.operators.extract.encrypted_documents import (
    DocumentDecryptor,
    DocumentPasswordsConfig,
    DocumentPasswordsConfigError,
)
from docpipe.core.operators.extract.ports.outbound.text_extraction import TextExtractionPort
from docpipe.exceptions.docpipe_exceptions import FlowExecutionFailedException

PASSWORD = "s3cret-Passw0rd"  # pragma: allowlist secret
WRONG_PASSWORD = "not-the-password"  # pragma: allowlist secret
ALL_PASSWORDS = (PASSWORD, WRONG_PASSWORD)


# ---------------------------------------------------------------------------
# Fixture builders
# ---------------------------------------------------------------------------


def _pdf(*, user_password: str | None = None, owner_password: str | None = None) -> bytes:
    """Build a one-page PDF, optionally encrypted with AES-256."""
    import pypdf

    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    if user_password is not None or owner_password is not None:
        writer.encrypt(user_password=user_password or "", owner_password=owner_password, algorithm="AES-256")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _encrypt_ooxml(*, data: bytes, password: str) -> bytes:
    from msoffcrypto.format.ooxml import OOXMLFile

    output = io.BytesIO()
    OOXMLFile(io.BytesIO(data)).encrypt(password, output)
    return output.getvalue()


def _docx(*, text: str = "quarterly figures", password: str | None = None) -> bytes:
    """Build a .docx with one paragraph, optionally encrypted."""
    import docx

    document = docx.Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    data = buffer.getvalue()
    return _encrypt_ooxml(data=data, password=password) if password else data


def _xlsx(*, password: str | None = None) -> bytes:
    """Build a .xlsx with one cell, optionally encrypted."""
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "budget"
    buffer = io.BytesIO()
    workbook.save(buffer)
    data = buffer.getvalue()
    return _encrypt_ooxml(data=data, password=password) if password else data


def _is_encrypted_pdf(data: bytes) -> bool:
    import pypdf

    return pypdf.PdfReader(io.BytesIO(data)).is_encrypted


def _decryptor(config: dict[str, Any] | None) -> DocumentDecryptor:
    return DocumentDecryptor(passwords_config=DocumentPasswordsConfig.from_config(config))


def _run_decrypt(*, decryptor: DocumentDecryptor, docs: list[tuple[str, str | None, bytes]]) -> list[dict[str, Any]]:
    """Run decrypt_tasks over (name, path, bytes) documents and return the tasks."""
    tasks = [
        {"idx": i, "doc_id": f"doc_{i}", "doc_name": name, OperatorConstants.Columns.BINARY_CONTENT: data}
        for i, (name, _, data) in enumerate(docs)
    ]
    columns: dict[str, list[Any]] = {"name": [name for name, _, _ in docs]}
    if any(path is not None for _, path, _ in docs):
        columns["path"] = [path for _, path, _ in docs]
    decryptor.decrypt_tasks(tasks=tasks, table=pa.table(columns))
    return tasks


def _failure_reason(task: dict[str, Any]) -> str | None:
    failed = task.get(DocumentDecryptor.FAILED_RESULT_KEY)
    return failed[OperatorConstants.Extraction.FAILURE_REASON] if failed else None


# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------


@pytest.mark.unit
def test_config_none_means_no_passwords() -> None:
    config = DocumentPasswordsConfig.from_config(None)
    assert config.rules == []
    assert config.default == []


@pytest.mark.unit
def test_config_valid_input_parses_and_hides_values() -> None:
    config = DocumentPasswordsConfig.from_config(
        {"rules": [{"match": "*.pdf", "passwords": [PASSWORD]}], "default": ["vault://org/shared-key"]}
    )
    assert config.rules[0].passwords[0].get_secret_value() == PASSWORD
    assert PASSWORD not in repr(config)


@pytest.mark.unit
@pytest.mark.parametrize(
    ("value", "expected_location"),
    [
        ({"unknown": [PASSWORD]}, "document_passwords.unknown"),
        ({"rules": [{"passwords": [PASSWORD]}]}, "document_passwords.rules[0].match"),
        ({"rules": [{"match": "", "passwords": [PASSWORD]}]}, "document_passwords.rules[0].match"),
        ({"rules": [{"match": "*.pdf", "passwords": []}]}, "document_passwords.rules[0].passwords"),
        ({"rules": [{"match": "*.pdf", "passwords": [PASSWORD, ""]}]}, "document_passwords.rules[0].passwords[1]"),
        (
            {"rules": [{"match": "*.pdf", "passwords": [PASSWORD], "note": PASSWORD}]},
            "document_passwords.rules[0].note",
        ),
        ({"default": [PASSWORD, 42]}, "document_passwords.default[1]"),
        ({"default": PASSWORD}, "document_passwords.default"),
        (PASSWORD, "document_passwords"),
    ],
)
def test_config_malformed_input_names_location_without_echoing_passwords(value: Any, expected_location: str) -> None:
    with pytest.raises(DocumentPasswordsConfigError) as exc_info:
        DocumentPasswordsConfig.from_config(value)

    assert any(message.startswith(f"{expected_location}:") for message in exc_info.value.messages)
    assert PASSWORD not in str(exc_info.value)
    assert exc_info.value.__cause__ is None


@pytest.mark.unit
def test_plain_text_password_locations_ignores_vault_references_and_placeholder() -> None:
    config = DocumentPasswordsConfig.from_config(
        {
            "rules": [
                {"match": "a*", "passwords": ["vault://org/a", PASSWORD]},
                {"match": "b*", "passwords": [DocpipeConstants.VAULT_VALIDATION_PLACEHOLDER]},
            ],
            "default": [WRONG_PASSWORD],
        }
    )
    assert config.plain_text_password_locations() == [
        "document_passwords.rules[0].passwords[1]",
        "document_passwords.default[0]",
    ]


@pytest.mark.unit
def test_candidates_first_matching_rule_wins() -> None:
    config = DocumentPasswordsConfig.from_config(
        {"rules": [{"match": "*.pdf", "passwords": ["a"]}, {"match": "report*", "passwords": ["b"]}]}
    )
    candidates = config.candidates_for(name="report.pdf", path=None)
    assert [c.get_secret_value() for c in candidates] == ["a"]


@pytest.mark.unit
def test_candidates_match_on_path() -> None:
    config = DocumentPasswordsConfig.from_config({"rules": [{"match": "*/finance/*", "passwords": ["a"]}]})
    candidates = config.candidates_for(name="report.pdf", path="s3://finance/q1/report.pdf")
    assert [c.get_secret_value() for c in candidates] == ["a"]


@pytest.mark.unit
def test_candidates_match_on_name_regardless_of_path() -> None:
    config = DocumentPasswordsConfig.from_config({"rules": [{"match": "budget_*.xlsx", "passwords": ["a"]}]})
    candidates = config.candidates_for(name="budget_2024.xlsx", path="/data/elsewhere/budget_2024.xlsx")
    assert [c.get_secret_value() for c in candidates] == ["a"]


@pytest.mark.unit
def test_candidates_same_name_in_different_folders() -> None:
    config = DocumentPasswordsConfig.from_config(
        {"rules": [{"match": "*/hr/*", "passwords": ["hr"]}, {"match": "*/finance/*", "passwords": ["fin"]}]}
    )
    hr = config.candidates_for(name="report.pdf", path="/data/hr/report.pdf")
    finance = config.candidates_for(name="report.pdf", path="/data/finance/report.pdf")
    assert [c.get_secret_value() for c in hr] == ["hr"]
    assert [c.get_secret_value() for c in finance] == ["fin"]


@pytest.mark.unit
def test_candidates_match_is_case_sensitive_whole_string() -> None:
    config = DocumentPasswordsConfig.from_config({"rules": [{"match": "*.PDF", "passwords": ["a"]}]})
    assert config.candidates_for(name="report.pdf", path=None) == []
    config = DocumentPasswordsConfig.from_config({"rules": [{"match": "finance/*", "passwords": ["a"]}]})
    assert config.candidates_for(name="x.pdf", path="/data/finance/x.pdf") == []


@pytest.mark.unit
def test_candidates_fall_back_to_default() -> None:
    config = DocumentPasswordsConfig.from_config(
        {"rules": [{"match": "*.docx", "passwords": ["a"]}], "default": ["shared"]}
    )
    candidates = config.candidates_for(name="report.pdf", path=None)
    assert [c.get_secret_value() for c in candidates] == ["shared"]


# ---------------------------------------------------------------------------
# PDF detection and decryption
# ---------------------------------------------------------------------------


@pytest.mark.unit
def test_pdf_decrypted_with_first_working_candidate() -> None:
    decryptor = _decryptor({"default": [WRONG_PASSWORD, PASSWORD]})
    original = _pdf(user_password=PASSWORD)
    tasks = _run_decrypt(decryptor=decryptor, docs=[("a.pdf", None, original)])

    decrypted = tasks[0][OperatorConstants.Columns.BINARY_CONTENT]
    assert _failure_reason(tasks[0]) is None
    assert decrypted != original
    assert not _is_encrypted_pdf(decrypted)


@pytest.mark.unit
def test_pdf_without_candidates_fails_with_password_required() -> None:
    tasks = _run_decrypt(decryptor=_decryptor(None), docs=[("a.pdf", None, _pdf(user_password=PASSWORD))])

    failed = tasks[0][DocumentDecryptor.FAILED_RESULT_KEY]
    assert failed[OperatorConstants.Extraction.SUCCESS] is False
    assert failed[OperatorConstants.Extraction.FAILURE_REASON] == OperatorConstants.Extraction.PASSWORD_REQUIRED
    assert failed[OperatorConstants.Extraction.ERROR].startswith(OperatorConstants.Extraction.PASSWORD_REQUIRED)


@pytest.mark.unit
def test_pdf_with_only_wrong_candidates_fails_with_password_incorrect() -> None:
    decryptor = _decryptor({"default": [WRONG_PASSWORD]})
    tasks = _run_decrypt(decryptor=decryptor, docs=[("a.pdf", None, _pdf(user_password=PASSWORD))])

    failed = tasks[0][DocumentDecryptor.FAILED_RESULT_KEY]
    assert failed[OperatorConstants.Extraction.FAILURE_REASON] == OperatorConstants.Extraction.PASSWORD_INCORRECT
    assert failed[OperatorConstants.Extraction.ERROR].startswith(OperatorConstants.Extraction.PASSWORD_INCORRECT)
    assert all(password not in failed[OperatorConstants.Extraction.ERROR] for password in ALL_PASSWORDS)


@pytest.mark.unit
def test_pdf_rule_matched_on_path_is_used() -> None:
    decryptor = _decryptor({"rules": [{"match": "*/finance/*", "passwords": [PASSWORD]}], "default": [WRONG_PASSWORD]})
    tasks = _run_decrypt(
        decryptor=decryptor,
        docs=[
            ("report.pdf", "/data/finance/report.pdf", _pdf(user_password=PASSWORD)),
            ("report.pdf", "/data/hr/report.pdf", _pdf(user_password=PASSWORD)),
        ],
    )
    assert _failure_reason(tasks[0]) is None
    assert _failure_reason(tasks[1]) == OperatorConstants.Extraction.PASSWORD_INCORRECT


@pytest.mark.unit
def test_owner_password_only_pdf_passes_through_unchanged() -> None:
    original = _pdf(owner_password=PASSWORD)
    tasks = _run_decrypt(decryptor=_decryptor(None), docs=[("a.pdf", None, original)])

    assert _failure_reason(tasks[0]) is None
    assert tasks[0][OperatorConstants.Columns.BINARY_CONTENT] is original


@pytest.mark.unit
def test_unencrypted_pdf_matched_by_rule_passes_through_unchanged() -> None:
    original = _pdf()
    tasks = _run_decrypt(decryptor=_decryptor({"default": [PASSWORD]}), docs=[("a.pdf", None, original)])

    assert _failure_reason(tasks[0]) is None
    assert tasks[0][OperatorConstants.Columns.BINARY_CONTENT] is original


@pytest.mark.unit
def test_encrypted_pdf_detected_by_header_without_pdf_extension() -> None:
    tasks = _run_decrypt(decryptor=_decryptor(None), docs=[("scan", None, _pdf(user_password=PASSWORD))])
    assert _failure_reason(tasks[0]) == OperatorConstants.Extraction.PASSWORD_REQUIRED


@pytest.mark.unit
def test_unreadable_pdf_passes_through_for_extraction_to_report() -> None:
    original = b"%PDF-1.7 /Encrypt this is not a real pdf"
    tasks = _run_decrypt(decryptor=_decryptor({"default": [PASSWORD]}), docs=[("a.pdf", None, original)])

    assert _failure_reason(tasks[0]) is None
    assert tasks[0][OperatorConstants.Columns.BINARY_CONTENT] is original


@pytest.mark.unit
def test_pdf_check_uses_docling_pdfium_lock() -> None:
    from docling.utils.locks import pypdfium2_lock

    assert DocumentDecryptor._pdfium_lock() is pypdfium2_lock


# ---------------------------------------------------------------------------
# OOXML detection and decryption
# ---------------------------------------------------------------------------


@pytest.mark.unit
def test_docx_decrypted_with_configured_password() -> None:
    import docx

    decryptor = _decryptor({"rules": [{"match": "*.docx", "passwords": [WRONG_PASSWORD, PASSWORD]}]})
    tasks = _run_decrypt(decryptor=decryptor, docs=[("w.docx", None, _docx(password=PASSWORD))])

    assert _failure_reason(tasks[0]) is None
    decrypted = tasks[0][OperatorConstants.Columns.BINARY_CONTENT]
    assert docx.Document(io.BytesIO(decrypted)).paragraphs[0].text == "quarterly figures"


@pytest.mark.unit
def test_xlsx_decrypted_with_configured_password() -> None:
    import openpyxl

    tasks = _run_decrypt(
        decryptor=_decryptor({"default": [PASSWORD]}), docs=[("b.xlsx", None, _xlsx(password=PASSWORD))]
    )

    assert _failure_reason(tasks[0]) is None
    decrypted = tasks[0][OperatorConstants.Columns.BINARY_CONTENT]
    assert openpyxl.load_workbook(io.BytesIO(decrypted)).active["A1"].value == "budget"


@pytest.mark.unit
def test_docx_without_candidates_fails_with_password_required() -> None:
    tasks = _run_decrypt(decryptor=_decryptor(None), docs=[("w.docx", None, _docx(password=PASSWORD))])
    assert _failure_reason(tasks[0]) == OperatorConstants.Extraction.PASSWORD_REQUIRED


@pytest.mark.unit
def test_docx_with_wrong_candidates_fails_with_password_incorrect() -> None:
    decryptor = _decryptor({"default": [WRONG_PASSWORD]})
    tasks = _run_decrypt(decryptor=decryptor, docs=[("w.docx", None, _docx(password=PASSWORD))])
    assert _failure_reason(tasks[0]) == OperatorConstants.Extraction.PASSWORD_INCORRECT


@pytest.mark.unit
def test_unencrypted_docx_passes_through_unchanged() -> None:
    original = _docx()
    tasks = _run_decrypt(decryptor=_decryptor({"default": [PASSWORD]}), docs=[("w.docx", None, original)])

    assert _failure_reason(tasks[0]) is None
    assert tasks[0][OperatorConstants.Columns.BINARY_CONTENT] is original


@pytest.mark.unit
def test_encrypted_docx_without_library_fails_with_decryption_unavailable(monkeypatch: pytest.MonkeyPatch) -> None:
    data = _docx(password=PASSWORD)
    monkeypatch.setitem(sys.modules, "msoffcrypto", None)
    monkeypatch.setitem(sys.modules, "msoffcrypto.exceptions", None)

    tasks = _run_decrypt(decryptor=_decryptor({"default": [PASSWORD]}), docs=[("w.docx", None, data)])

    failed = tasks[0][DocumentDecryptor.FAILED_RESULT_KEY]
    assert failed[OperatorConstants.Extraction.FAILURE_REASON] == OperatorConstants.Extraction.DECRYPTION_UNAVAILABLE
    assert "docling-pipelines-slim[extract]" in failed[OperatorConstants.Extraction.ERROR]


@pytest.mark.unit
def test_ole_file_without_encrypted_package_passes_through_without_library(monkeypatch: pytest.MonkeyPatch) -> None:
    original = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512
    monkeypatch.setitem(sys.modules, "msoffcrypto", None)
    monkeypatch.setitem(sys.modules, "msoffcrypto.exceptions", None)

    tasks = _run_decrypt(decryptor=_decryptor(None), docs=[("legacy.docx", None, original)])

    assert _failure_reason(tasks[0]) is None
    assert tasks[0][OperatorConstants.Columns.BINARY_CONTENT] is original


@pytest.mark.unit
def test_unsupported_office_encryption_scheme_fails_with_decryption_unavailable() -> None:
    import msoffcrypto
    from msoffcrypto.exceptions import DecryptionError

    data = _docx(password=PASSWORD)
    with patch.object(msoffcrypto, "OfficeFile", side_effect=DecryptionError("Unsupported EncryptionInfo version")):
        tasks = _run_decrypt(decryptor=_decryptor({"default": [PASSWORD]}), docs=[("w.docx", None, data)])

    assert _failure_reason(tasks[0]) == OperatorConstants.Extraction.DECRYPTION_UNAVAILABLE


@pytest.mark.unit
@pytest.mark.parametrize("name", ["image.png", "notes.html", "audio.mp3", "data.csv"])
def test_unsupported_formats_pass_through_unchanged(name: str) -> None:
    original = b"/Encrypt \xd0\xcf\x11\xe0 not a pdf"
    tasks = _run_decrypt(decryptor=_decryptor({"default": [PASSWORD]}), docs=[(name, None, original)])

    assert _failure_reason(tasks[0]) is None
    assert tasks[0][OperatorConstants.Columns.BINARY_CONTENT] is original


@pytest.mark.unit
def test_tasks_with_fetch_errors_are_left_alone() -> None:
    task = {"idx": 0, "doc_id": "doc_0", "doc_name": "a.pdf", "error": "fetch failed"}
    _decryptor(None).decrypt_tasks(tasks=[task], table=pa.table({"name": ["a.pdf"]}))
    assert task == {"idx": 0, "doc_id": "doc_0", "doc_name": "a.pdf", "error": "fetch failed"}


@pytest.mark.unit
def test_failed_result_future_only_for_failed_tasks() -> None:
    assert DocumentDecryptor.failed_result_future(task={"idx": 0}) is None

    failed = {OperatorConstants.Extraction.SUCCESS: False}
    future = DocumentDecryptor.failed_result_future(task={DocumentDecryptor.FAILED_RESULT_KEY: failed})
    assert future is not None
    assert future.done()
    assert future.result() is failed


@pytest.mark.unit
def test_decryption_never_logs_passwords(caplog: pytest.LogCaptureFixture) -> None:
    decryptor = _decryptor({"rules": [{"match": "*", "passwords": [WRONG_PASSWORD, PASSWORD]}]})
    handler_records: list[logging.LogRecord] = []

    class _Collect(logging.Handler):
        def emit(self, record: logging.LogRecord) -> None:
            handler_records.append(record)

    handler = _Collect(level=logging.DEBUG)
    encrypted_documents.logger.addHandler(handler)
    try:
        _run_decrypt(
            decryptor=decryptor,
            docs=[
                ("a.pdf", None, _pdf(user_password=PASSWORD)),
                ("b.pdf", None, _pdf(user_password="other")),  # pragma: allowlist secret
                ("w.docx", None, _docx(password=PASSWORD)),
            ],
        )
    finally:
        encrypted_documents.logger.removeHandler(handler)

    assert handler_records
    logged = " ".join(record.getMessage() for record in handler_records)
    assert all(password not in logged for password in ALL_PASSWORDS)


# ---------------------------------------------------------------------------
# Sequential path: TextExtractionPort
# ---------------------------------------------------------------------------


class _RecordingAdapter(TextExtractionPort):
    """Text extraction adapter that records the bytes it receives."""

    ADAPTER_NAME = "recording"
    ADAPTER_DISPLAY_NAME = "Recording Adapter"

    def __init__(self, *, config: dict[str, Any]) -> None:
        super().__init__(config=config)
        self.received: dict[str, bytes] = {}

    @staticmethod
    def get_config_schema() -> type[BaseModel]:
        return BaseModel

    def extract_single_document(self, *, file_path: str, binary_content: bytes, **kwargs: Any) -> dict[str, Any]:
        self.received[file_path] = binary_content
        return {
            OperatorConstants.Extraction.SUCCESS: True,
            OperatorConstants.Columns.DOC_COLUMN_DEFAULT: f"text of {file_path}",
            OperatorConstants.Metadata.METADATA: {"page_count": 1},
        }


def _mixed_table() -> pa.Table:
    """One unencrypted PDF, one encrypted PDF with a configured password, one without."""
    return pa.table(
        {
            "id": ["plain", "known", "unknown"],
            "name": ["plain.pdf", "known.pdf", "unknown.pdf"],
            "path": ["/data/plain.pdf", "/data/finance/known.pdf", "/data/hr/unknown.pdf"],
            "binary_content": [_pdf(), _pdf(user_password=PASSWORD), _pdf(user_password=PASSWORD)],
        }
    )


_FINANCE_ONLY = {"rules": [{"match": "*/finance/*", "passwords": [WRONG_PASSWORD, PASSWORD]}]}


@pytest.mark.unit
def test_port_mixed_batch_decrypts_and_reports_recoverable_failure() -> None:
    adapter = _RecordingAdapter(config={"doc_column": "content", "max_workers": 2})
    adapter.document_decryptor = _decryptor(_FINANCE_ONLY)
    table = _mixed_table()
    original_bytes = table.column("binary_content").to_pylist()
    metadata = {
        Metrics.External.PROCESSED_DOCS: 0,
        Metrics.External.FAILED_DOCS: [],
        Metrics.External.FAILED_DOCS_COUNT: 0,
    }

    result_tables, result_meta = adapter.transform(table=table, metadata=metadata)

    output = result_tables[0]
    assert output.num_rows == 2
    assert sorted(output.column("name").to_pylist()) == ["known.pdf", "plain.pdf"]
    assert result_meta[Metrics.External.PROCESSED_DOCS] == 2
    assert result_meta[Metrics.External.FAILED_DOCS_COUNT] == 1
    failed = result_meta[Metrics.External.FAILED_DOCS][0]
    assert failed["name"] == "unknown.pdf"
    assert failed["reason"].startswith(OperatorConstants.Extraction.PASSWORD_REQUIRED)
    assert "future flow executions" not in failed["reason"]
    assert Metrics.Internal.NON_RECOVERABLE_DOCS_TABLE not in result_meta

    # The adapter got plain bytes for the decrypted document and never saw the failed one
    assert not _is_encrypted_pdf(adapter.received["known.pdf"])
    assert adapter.received["plain.pdf"] == original_bytes[0]
    assert "unknown.pdf" not in adapter.received

    # The input table is untouched and nothing secret reaches the outputs
    assert table.column("binary_content").to_pylist() == original_bytes
    rendered = repr(result_meta) + repr(output.to_pylist())
    assert all(password not in rendered for password in ALL_PASSWORDS)


@pytest.mark.unit
def test_port_without_decryptor_passes_bytes_through() -> None:
    adapter = _RecordingAdapter(config={"doc_column": "content", "max_workers": 1})
    encrypted = _pdf(user_password=PASSWORD)
    table = pa.table({"id": ["a"], "name": ["a.pdf"], "binary_content": [encrypted]})
    metadata = {
        Metrics.External.PROCESSED_DOCS: 0,
        Metrics.External.FAILED_DOCS: [],
        Metrics.External.FAILED_DOCS_COUNT: 0,
    }

    adapter.transform(table=table, metadata=metadata)

    assert adapter.received["a.pdf"] == encrypted


# ---------------------------------------------------------------------------
# ExtractOperator: configuration, validation, metadata and streaming path
# ---------------------------------------------------------------------------


def _operator_config(**extra: Any) -> dict[str, Any]:
    return {"text_extraction": {"provider": "docling_library"}, **extra}


@pytest.mark.unit
def test_operator_attaches_decryptor_to_text_adapter() -> None:
    from docpipe.core.operators.extract.extract_operator import ExtractOperator

    operator = ExtractOperator(config=_operator_config(document_passwords={"default": [PASSWORD]}))

    assert operator.text_adapter.document_decryptor is operator.document_decryptor
    assert operator.document_passwords_config.default[0].get_secret_value() == PASSWORD


@pytest.mark.unit
def test_operator_detects_encryption_without_configuration() -> None:
    from docpipe.core.operators.extract.extract_operator import ExtractOperator

    operator = ExtractOperator(config=_operator_config())

    assert operator.text_adapter.document_decryptor is operator.document_decryptor
    assert operator.document_passwords_config.rules == []


@pytest.mark.unit
def test_operator_rejects_malformed_passwords_without_echoing_them() -> None:
    from docpipe.core.operators.extract.extract_operator import ExtractOperator

    with pytest.raises(FlowExecutionFailedException) as exc_info:
        ExtractOperator(config=_operator_config(document_passwords={"rules": [{"passwords": [PASSWORD]}]}))

    assert "document_passwords.rules[0].match" in str(exc_info.value)
    assert PASSWORD not in str(exc_info.value)


@pytest.mark.unit
def test_operator_validate_warns_on_plain_text_passwords() -> None:
    from docpipe.core.operators.extract.extract_operator import ExtractOperator

    operator = ExtractOperator(
        config=_operator_config(
            document_passwords={
                "rules": [{"match": "*.pdf", "passwords": [DocpipeConstants.VAULT_VALIDATION_PLACEHOLDER]}],
                "default": [PASSWORD],
            }
        )
    )
    errors: list[Any] = []
    warnings: list[Any] = []
    operator.validate(errors, warnings, ["content"])

    password_warnings = [w for w in warnings if "plain-text password" in str(w)]
    assert len(password_warnings) == 1
    assert "document_passwords.default[0]" in str(password_warnings[0])
    assert PASSWORD not in str(warnings)
    assert errors == []


@pytest.mark.unit
def test_operator_validate_no_warning_for_vault_references() -> None:
    from docpipe.core.operators.extract.extract_operator import ExtractOperator

    operator = ExtractOperator(
        config=_operator_config(document_passwords={"default": ["vault://org/key"]}),
    )
    warnings: list[Any] = []
    operator.validate([], warnings, ["content"])

    assert not any("plain-text password" in str(w) for w in warnings)


@pytest.mark.unit
def test_flow_validator_placeholder_matches_shared_constant() -> None:
    from docpipe.core.orchestration.flow_validator import FlowValidator

    sanitized = FlowValidator._sanitize_vault_refs_for_validation({"default": ["vault://org/key"]})
    assert sanitized == {"default": [DocpipeConstants.VAULT_VALIDATION_PLACEHOLDER]}


@pytest.mark.unit
def test_operator_metadata_marks_document_passwords_sensitive() -> None:
    from docpipe.core.operators.extract.extract_operator import ExtractOperator

    attributes = ExtractOperator.get_metadata()[OperatorConstants.Config.ATTRIBUTES]
    attribute = attributes[OperatorConstants.Config.DOCUMENT_PASSWORDS]

    assert attribute[OperatorConstants.Config.REQUIRED] is False
    assert attribute[OperatorConstants.Config.SENSITIVE] is True


def _streaming_operator(*, document_passwords: dict[str, Any] | None) -> Any:
    """ExtractOperator with a mocked LiteLLM entity adapter that reads document bytes."""
    from docpipe.core.operators.extract.extract_operator import ExtractOperator
    from docpipe.integrations.litellm.client import LiteLLMLLMClient

    config = _operator_config(
        entity_extraction={
            "provider": "litellm",
            "provider_config": {
                "model_id": "openai/llama3.2",
                "api_base": "http://localhost:11434/v1",
                "api_key": "test",  # pragma: allowlist secret
            },
            "custom_schema": {"title": "string"},
        },
        max_workers=2,
    )
    if document_passwords is not None:
        config[OperatorConstants.Config.DOCUMENT_PASSWORDS] = document_passwords
    with patch.object(LiteLLMLLMClient, "__init__", return_value=None):
        operator = ExtractOperator(config=config)

    received: dict[str, bytes] = {}

    def _extract(*, file_path: str, binary_content: bytes, **kwargs: Any) -> dict[str, Any]:
        received[file_path] = binary_content
        return {"success": True, "content": f"text of {file_path}", "metadata": {"page_count": 1}}

    operator.text_adapter.extract_single_document = MagicMock(side_effect=_extract)
    operator.entity_adapter.requires_binary_content = True
    operator.entity_adapter.supported_extensions = frozenset({".pdf"})
    operator.entity_adapter.extract_entities_single = MagicMock(
        return_value={"success": True, "entities": {"title": "t"}, "error": None}
    )
    return operator, received


@pytest.mark.unit
def test_streaming_path_decrypts_for_text_and_binary_entity_adapters() -> None:
    operator, received = _streaming_operator(document_passwords=_FINANCE_ONLY)
    table = _mixed_table()
    metadata = operator.create_base_metadata(total_docs_count=table.num_rows)

    result_tables, result_meta = operator._run_streaming_pipeline(table=table, metadata=metadata)

    assert result_tables[0].num_rows == 2
    assert result_meta[Metrics.External.PROCESSED_DOCS] == 2
    assert result_meta[Metrics.External.FAILED_DOCS_COUNT] == 1
    failed = result_meta[Metrics.External.FAILED_DOCS][0]
    assert failed["name"] == "unknown.pdf"
    assert failed["reason"].startswith(OperatorConstants.Extraction.PASSWORD_REQUIRED)
    assert Metrics.Internal.NON_RECOVERABLE_DOCS_TABLE not in result_meta

    assert not _is_encrypted_pdf(received["known.pdf"])
    assert "unknown.pdf" not in received
    entity_bytes = {
        call.kwargs["doc_name"]: call.kwargs["content"]
        for call in operator.entity_adapter.extract_entities_single.call_args_list
    }
    assert entity_bytes["known.pdf"] == received["known.pdf"]

    rendered = repr(result_meta) + repr(result_tables[0].to_pylist())
    assert all(password not in rendered for password in ALL_PASSWORDS)


@pytest.mark.unit
def test_streaming_path_reports_password_incorrect_as_recoverable() -> None:
    operator, _ = _streaming_operator(document_passwords={"default": [WRONG_PASSWORD]})
    table = pa.table(
        {
            "id": ["plain", "locked"],
            "name": ["plain.pdf", "locked.pdf"],
            "path": ["/data/plain.pdf", "/data/locked.pdf"],
            "binary_content": [_pdf(), _pdf(user_password=PASSWORD)],
        }
    )
    metadata = operator.create_base_metadata(total_docs_count=table.num_rows)

    _, result_meta = operator._run_streaming_pipeline(table=table, metadata=metadata)

    failed = result_meta[Metrics.External.FAILED_DOCS][0]
    assert failed["reason"].startswith(OperatorConstants.Extraction.PASSWORD_INCORRECT)
    assert "future flow executions" not in failed["reason"]
    assert Metrics.Internal.NON_RECOVERABLE_DOCS_TABLE not in result_meta
